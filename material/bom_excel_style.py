"""Presentation only for the existing 13-column structured BOM export."""
from functools import lru_cache

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def format_structured_bom(ws):
    """Keep cell values, coordinates, row order and blank cells unchanged."""
    navy, blue, muted = '243B53', '334E68', '627D98'
    fonts = {
        'header': Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF'),
        'root': Font(name='맑은 고딕', size=10, bold=True, color=navy),
        'semi': Font(name='맑은 고딕', size=10, bold=True, color=blue),
        'body': Font(name='맑은 고딕', size=10, color='243746'),
        'muted': Font(name='맑은 고딕', size=10, color=muted),
        'shortage': Font(name='맑은 고딕', size=10, bold=True, color='B42318'),
        'missing': Font(name='맑은 고딕', size=10, bold=True, color='92400E'),
    }
    fills = {name: PatternFill('solid', fgColor=color) for name, color in {
        'header': navy, 'root': 'DCEAF2', 'semi': 'EDF3F8',
        'white': 'FFFFFF', 'stripe': 'F8FAFC', 'shortage': 'FEE4E2', 'missing': 'FFF4D6',
    }.items()}
    thin = Side(style='hair', color='DDE5ED')
    group_edge = Side(style='medium', color='829AB1')
    borders = {
        'body': Border(bottom=thin),
        'root': Border(top=group_edge, bottom=thin),
        'header': Border(bottom=Side(style='medium', color='486581')),
    }

    @lru_cache(maxsize=32)
    def alignment(horizontal, indent=0, wrap=False):
        return Alignment(horizontal=horizontal, vertical='center', indent=indent, wrap_text=wrap)

    widths = [19, 32, 13, 15, 8, 25, 46, 9, 15, 16, 18, 16, 34]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.font = fonts['header']
        cell.fill = fills['header']
        cell.alignment = alignment('center', wrap=True)
        cell.border = borders['header']

    for row in ws.iter_rows(min_row=2, max_col=13):
        # The exporter already marks semi-finished rows in blue. Read that marker
        # before replacing styles, so leaf items with deeper LEVEL remain leaves.
        level = row[4].value
        root = level == 0 or row[0].value not in (None, '')
        semi = row[5].fill.fgColor.rgb == '00D9E2F3'
        missing = row[5].value == 'BOM 없음'
        kind = 'missing' if missing else 'root' if root else 'semi' if semi else 'body'
        fill = fills[kind] if kind != 'body' else fills['stripe' if row[0].row % 2 else 'white']
        ws.row_dimensions[row[0].row].height = 34 if root or missing else 30
        for index, cell in enumerate(row, 1):
            cell.font = fonts[kind]
            cell.fill = fill
            cell.border = borders['root' if root else 'body']
            if index in (3, 9, 10, 11, 12):
                cell.alignment = alignment('right')
                # Optional decimals retain small BOM quantities; numeric values
                # are never rounded or converted to strings by this formatter.
                cell.number_format = '#,##0.###############;[Red]-#,##0.###############;0'
                if isinstance(cell.value, (int, float)):
                    if cell.value == int(cell.value):
                        cell.number_format = '#,##0;[Red]-#,##0;0'
                    elif abs(cell.value) < 1e-15:
                        cell.number_format = '0.##############E+00'
            elif index in (4, 5, 8):
                cell.alignment = alignment('center')
            elif index == 7:
                indent = min(max(int(level or 0) - 1, 0), 6)
                cell.alignment = alignment('left', indent, True)
            else:
                cell.alignment = alignment('left', wrap=index in (2, 13))
            if kind == 'body' and index in (4, 5, 8):
                cell.font = fonts['muted']
        shortage = row[11]
        if isinstance(shortage.value, (int, float)):
            if shortage.value > 0:
                shortage.fill = fills['shortage']
                shortage.font = fonts['shortage']
            elif shortage.value == 0:
                shortage.font = fonts['muted']

    ws.freeze_panes = 'F2'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.print_title_rows = '1:1'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    ws.oddFooter.center.text = '&P / &N'
