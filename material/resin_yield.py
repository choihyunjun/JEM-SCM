"""성형 원단위 중량 관리: 월별 품번·재질·원재료별 수지 사용량과 스크랩.

출처
- 생산수(양품/불량), 파지 중량: ERP 생산실적 API. 파지는 비고(remarkDc)의 숫자 g (M7번대 4자리, 그 외 5자리)
- 제품중량·런너중량·C/V·총소요량·재료: 성형 마스터 (품번 × 재료구분 행)
- 검사불량: 불량창고로 옮긴 이동 수량 (ERP 재고이동 동기화 + SCM 이동)

계산 (재료 행마다, kg)
- 양품/공정불량/검사불량 = 수량 × 제품중량, 런너 = 생산수 × 런너중량 ÷ C/V
- 파지 = 품번의 비고 합계를 재료 행의 총소요량 비율로 나눔
- 총 사용량 = 양품 + 불량 + 런너 + 파지, 스크랩 = 불량 + 런너 + 파지
"""
import re
from collections import defaultdict
from datetime import date, datetime

from django.conf import settings
from django.core.cache import cache
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from orders.decorators import wms_permission_required
from .erp_api import call_erp_api
from .models import MaterialTransaction, MoldingMaster, ResinType

WR_API = '/apiproxy/api20A03S00901'
CACHE_SECONDS = 30 * 60
SCRAP_RE = re.compile(r'^(?:[A-Za-z]\d{3}\s+)?(\d{4,5})$')  # '01475', '6100', 과거 엑셀 형식 'A670 01036'
KNOWN_RESINS = {'ABS', 'ASA', 'PA', 'PA6', 'PA12', 'PA46', 'PA66', 'PBT', 'PC', 'PC+ABS', 'PMMA', 'POM', 'PP',
                'TPE', 'TPO', 'TPU', 'M/B'}
MATERIAL_ORDER = {'재료1': 0, '재료2': 1, 'M/B': 2}


def parse_scrap(remark, equip):
    """비고에서 파지 중량(g)을 읽는다. (g, 확인 사유 또는 None)"""
    text = str(remark or '').strip()
    if not text:
        return 0, None
    m = SCRAP_RE.match(text)
    if not m:
        return 0, '숫자 형식 아님'
    digits = m.group(1)
    expected = 4 if equip.startswith('M7') else 5
    issue = None if len(digits) == expected else f'{len(digits)}자리 (설비 {equip}는 {expected}자리)'
    return int(digits), issue


def month_range(year, month):
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return start, end


def fetch_results(year, month, refresh=False):
    """ERP 생산실적(성형 설비 M*)을 월 단위로 조회. 30분 캐시. ({'rows', 'fetched_at'}, 오류)"""
    key = f'resin_yield:wr:{year}{month:02d}'
    if not refresh:
        cached = cache.get(key)
        if cached is not None:
            return cached, None
    start, end = month_range(year, month)
    last = date.fromordinal(end.toordinal() - 1)
    ok, raw, err = call_erp_api(WR_API, {'coCd': settings.ERP_COMPANY_CODE,
                                         'wrDtFrom': start.strftime('%Y%m%d'), 'wrDtTo': last.strftime('%Y%m%d')})
    if not ok:
        return None, err or 'ERP 생산실적 조회 실패'
    rows = []
    for r in (raw or {}).get('resultData', []) or []:
        equip = (r.get('equipNm') or '').strip()
        part_no = (r.get('itemCd') or '').strip()
        if not re.match(r'^M\d', equip) or not part_no:
            continue
        rows.append({
            'date': r.get('wrDt', ''), 'equip': equip, 'part_no': part_no,
            'part_name': (r.get('itemNm') or '').strip(), 'wr_no': r.get('wrCd', ''),
            'good': float(r.get('goodQt') or 0), 'bad': float(r.get('badQt') or 0),
            'remark': (r.get('remarkDc') or '').strip(),
        })
    data = {'rows': rows, 'fetched_at': timezone.localtime().strftime('%Y-%m-%d %H:%M')}
    cache.set(key, data, CACHE_SECONDS)
    return data, None


def inspection_defects(year, month):
    """불량창고로 옮긴 품번별 수량 (재고 맞춤 기록은 제외)."""
    start, end = month_range(year, month)
    tz = timezone.get_current_timezone()
    qs = MaterialTransaction.objects.filter(
        transaction_type__in=('TRF_ERP', 'TRANSFER'), warehouse_to__name__contains='불량',
        date__gte=datetime(start.year, start.month, 1, tzinfo=tz),
        date__lt=datetime(end.year, end.month, 1, tzinfo=tz),
    ).values('part__part_no').annotate(qty=Sum('quantity'))
    return {row['part__part_no']: float(row['qty'] or 0) for row in qs}


def resin_of(line, mapping):
    if line.material_part_no in mapping:
        return mapping[line.material_part_no]
    if line.material_type == 'M/B':
        return 'M/B'
    first = (line.material_name or '').split(' ')[0].upper()
    first = 'M/B' if first == 'MB' else first
    return first if first in KNOWN_RESINS else '미분류'


def _ratio(part, whole):
    return part / whole if whole else 0


def build_report(year, month, refresh=False):
    data, err = fetch_results(year, month, refresh)
    if err:
        return {'error': err}

    produced = defaultdict(lambda: {'good': 0.0, 'bad': 0.0, 'scrap_g': 0, 'name': ''})
    issues = []
    for r in data['rows']:
        p = produced[r['part_no']]
        p['good'] += r['good']
        p['bad'] += r['bad']
        p['name'] = p['name'] or r['part_name']
        grams, issue = parse_scrap(r['remark'], r['equip'])
        p['scrap_g'] += grams
        if issue:
            issues.append({**r, 'issue': issue, 'grams': grams})

    defects = inspection_defects(year, month)
    part_nos = set(produced) | {k for k, v in defects.items() if v}
    masters = defaultdict(list)
    for line in MoldingMaster.objects.filter(part_no__in=part_nos):
        masters[line.part_no].append(line)
    mapping = dict(ResinType.objects.values_list('material_part_no', 'resin_type'))

    lines, unregistered = [], []
    for part_no in sorted(part_nos):
        p = produced.get(part_no) or {'good': 0.0, 'bad': 0.0, 'scrap_g': 0, 'name': ''}
        made = p['good'] + p['bad']
        defect_qty = defects.get(part_no, 0.0)
        part_lines = sorted(masters.get(part_no, []), key=lambda m: (MATERIAL_ORDER.get(m.material_type, 9), m.pk))
        if not part_lines:
            if made or p['scrap_g']:  # 불량창고 이동만 있는 조립품 등 성형품이 아닌 품번은 무시
                unregistered.append({'part_no': part_no, 'part_name': p['name'], 'produced': made,
                                     'scrap_kg': p['scrap_g'] / 1000, 'defect_qty': defect_qty})
            continue
        total_material = sum(float(m.total_material) for m in part_lines)
        for m in part_lines:
            weight = float(m.product_weight)
            runner = float(m.runner_weight)
            cavity = m.cavity or 0
            share = _ratio(float(m.total_material), total_material) if total_material else 1 / len(part_lines)
            row = {
                'part_no': part_no, 'part_name': m.part_name or p['name'], 'resin': resin_of(m, mapping),
                'material_type': m.material_type, 'material_part_no': m.material_part_no,
                'material_name': m.material_name, 'cavity': cavity, 'product_weight': weight,
                'runner_weight': runner,
                'unit_standard': weight + (runner / cavity if cavity else 0),
                'produced': made, 'good_qty': p['good'], 'bad_qty': p['bad'], 'defect_qty': defect_qty,
                'good_kg': p['good'] * weight / 1000,
                'bad_kg': p['bad'] * weight / 1000,
                'inspect_kg': defect_qty * weight / 1000,
                'runner_kg': made * runner / cavity / 1000 if cavity else 0,
                'scrap_kg': p['scrap_g'] / 1000 * share,
                'no_cavity': not cavity,
            }
            row['loss_kg'] = row['bad_kg'] + row['inspect_kg'] + row['runner_kg'] + row['scrap_kg']
            row['total_kg'] = row['good_kg'] + row['loss_kg']
            row['loss_rate'] = _ratio(row['loss_kg'], row['total_kg'])
            row['loss_pct'] = row['loss_rate'] * 100
            row['unit_actual'] = row['total_kg'] * 1000 / made if made else None
            lines.append(row)

    return {
        'lines': lines,
        'by_resin': _summarize(lines, 'resin'),
        'by_material': _summarize(lines, 'material_part_no'),
        'totals': _summarize(lines, None)[0] if lines else None,
        'issues': issues,
        'unregistered': unregistered,
        'unregistered_produced': sum(u['produced'] for u in unregistered),
        'unregistered_scrap_kg': sum(u['scrap_kg'] for u in unregistered),
        'no_cavity': [r for r in lines if r['no_cavity'] and r['produced']],
        'fetched_at': data['fetched_at'],
        'erp_rows': len(data['rows']),
    }


KG_FIELDS = ('good_kg', 'bad_kg', 'inspect_kg', 'runner_kg', 'scrap_kg', 'loss_kg', 'total_kg')


def _summarize(lines, key):
    groups = {}
    for row in lines:
        k = row[key] if key else '합계'
        g = groups.setdefault(k, {'key': k, 'parts': set(), 'produced': 0.0, **{f: 0.0 for f in KG_FIELDS},
                                  'material_name': row['material_name'], 'resin': row['resin']})
        if row['part_no'] not in g['parts']:  # 재료가 여러 행인 품번의 생산수는 한 번만
            g['parts'].add(row['part_no'])
            g['produced'] += row['produced']
        for f in KG_FIELDS:
            g[f] += row[f]
    result = []
    for g in sorted(groups.values(), key=lambda g: -g['total_kg']):
        g['part_count'] = len(g.pop('parts'))
        g['loss_rate'] = _ratio(g['loss_kg'], g['total_kg'])
        g['loss_pct'] = g['loss_rate'] * 100
        result.append(g)
    return result


def _month_param(request):
    raw = (request.GET.get('month') or '').strip()
    try:
        picked = datetime.strptime(raw, '%Y-%m')
        return picked.year, picked.month
    except ValueError:
        today = timezone.localdate()
        return today.year, today.month


@wms_permission_required('can_wms_stock_view')
def resin_yield(request):
    year, month = _month_param(request)
    report = build_report(year, month, refresh=request.GET.get('refresh') == '1')
    return render(request, 'material/resin_yield.html', {
        'report': report, 'month': f'{year}-{month:02d}', 'year': year, 'mon': month,
        'resins': sorted({r['resin'] for r in report.get('lines', [])}),
        'issue_count': len(report.get('issues', [])) + len(report.get('unregistered', [])),
    })


@wms_permission_required('can_wms_stock_view')
def resin_yield_excel(request):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    year, month = _month_param(request)
    report = build_report(year, month)
    if report.get('error'):
        return HttpResponse(f"ERP 생산실적 조회 실패: {report['error']}", status=502,
                            content_type='text/plain; charset=utf-8')

    wb = openpyxl.Workbook()
    fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

    def sheet(ws, title, headers, rows, formats):
        ws.title = title
        ws.append([f'{title} ({year}년 {month}월)'])
        ws['A1'].font = Font(bold=True, size=12)
        ws.append(headers)
        for cell in ws[2]:
            cell.fill = fill
            cell.font = Font(bold=True, color='FFFFFF', size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in rows:
            ws.append(row)
        for col, fmt in formats.items():
            for cells in ws.iter_rows(min_row=3, min_col=col, max_col=col):
                cells[0].number_format = fmt
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(32, len(str(h)) * 2 + 2))
        ws.freeze_panes = 'A3'

    kg, g, pct, ea = '#,##0.000', '#,##0.000', '0.0%', '#,##0'
    lines = report['lines']
    sheet(wb.active, '품번별 원단위', [
        'NO', '품번', '품명', '재질', '재료구분', '재료품번', '재료명', 'C/V', '제품중량(g/EA)', '런너중량(g/SHOT)',
        '기준 원단위(g/EA)', '생산수(EA)', '양품량(kg)', '공정불량(kg)', '검사불량(kg)', '런너량(kg)', '파지량(kg)',
        '총 수지사용량(kg)', '스크랩량(kg)', '스크랩률', '실적 원단위(g/EA)',
    ], [[i, r['part_no'], r['part_name'], r['resin'], r['material_type'], r['material_part_no'], r['material_name'],
         r['cavity'], r['product_weight'], r['runner_weight'], r['unit_standard'], r['produced'], r['good_kg'],
         r['bad_kg'], r['inspect_kg'], r['runner_kg'], r['scrap_kg'], r['total_kg'], r['loss_kg'], r['loss_rate'],
         r['unit_actual']] for i, r in enumerate(lines, 1)],
        {9: g, 10: g, 11: g, 12: ea, 13: kg, 14: kg, 15: kg, 16: kg, 17: kg, 18: kg, 19: kg, 20: pct, 21: g})

    summary_headers = ['생산수(EA)', '양품량(kg)', '공정불량(kg)', '검사불량(kg)', '런너량(kg)', '파지량(kg)',
                       '총 수지사용량(kg)', '스크랩량(kg)', '스크랩률']

    def summary_values(s):
        return [s['produced'], s['good_kg'], s['bad_kg'], s['inspect_kg'], s['runner_kg'], s['scrap_kg'],
                s['total_kg'], s['loss_kg'], s['loss_rate']]

    totals = [report['totals']] if report['totals'] else []
    sheet(wb.create_sheet(), '재질별 집계', ['재질', '품번 수'] + summary_headers,
          [[s['key'], s['part_count']] + summary_values(s) for s in report['by_resin']]
          + [['합계', t['part_count']] + summary_values(t) for t in totals],
          {3: ea, 4: kg, 5: kg, 6: kg, 7: kg, 8: kg, 9: kg, 10: kg, 11: pct})
    sheet(wb.create_sheet(), '원재료별 사용량', ['재료품번', '재료명', '재질', '품번 수'] + summary_headers,
          [[s['key'], s['material_name'], s['resin'], s['part_count']] + summary_values(s)
           for s in report['by_material']],
          {5: ea, 6: kg, 7: kg, 8: kg, 9: kg, 10: kg, 11: kg, 12: kg, 13: pct})
    sheet(wb.create_sheet(), '확인 필요', ['구분', '품번', '품명', '설비', '실적일', '실적번호', '비고/내용', '반영'],
          [['파지 비고 형식', i['part_no'], i['part_name'], i['equip'], i['date'], i['wr_no'], i['remark'],
            f"{i['issue']} → {i['grams']:,}g" if i['grams'] else f"{i['issue']} → 0g"] for i in report['issues']]
          + [['성형 마스터 미등록', u['part_no'], u['part_name'], '', '', '',
              f"생산 {u['produced']:,.0f}EA / 파지 {u['scrap_kg']:,.3f}kg / 검사불량 {u['defect_qty']:,.0f}EA",
              '집계 제외'] for u in report['unregistered']]
          + [['C/V 미입력', r['part_no'], r['part_name'], '', '', '', r['material_part_no'], '런너량 0으로 계산']
             for r in report['no_cavity']], {})

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''resin_yield_{year}{month:02d}.xlsx"
    wb.save(response)
    return response
