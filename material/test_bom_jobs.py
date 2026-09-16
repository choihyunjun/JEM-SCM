import io
from datetime import date, timedelta
from unittest.mock import patch

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from orders.models import Demand, Part, Vendor
from .bom_jobs import advance_job, build_snapshot, create_job, parse_plans
from .models import BOMCalculationJob, BOMCalculationRow, BOMItem, MaterialStock, Product, Warehouse
from .views import _calculate_bom_requirements


def csv_file(rows):
    return SimpleUploadedFile('plan.csv', ('날짜,품번,계획수량\n' + rows).encode('utf-8-sig'))


class BOMJobTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='bom-admin', is_superuser=True)
        self.client.force_login(self.user)
        self.vendor = Vendor.objects.create(name='BOM vendor', code='BOM-V')
        self.part = Part.objects.create(part_no='RAW', part_name='Raw material', vendor=self.vendor)
        self.product = Product.objects.create(part_no='FG', part_name='Finished')
        BOMItem.objects.create(product=self.product, seq=1, child_part_no='RAW', child_part_name='Raw material', required_qty=1)
        wh = Warehouse.objects.create(code='BOM-WH', name='BOM warehouse')
        self.stock = MaterialStock.objects.create(part=self.part, warehouse=wh, quantity=100)

    def job(self, rows='20260901,FG,80\n20260902,FG,80'):
        return create_job(self.user, csv_file(rows))

    def finish(self, job):
        for _ in range(100):
            advance_job(job)
            job.refresh_from_db()
            if job.status in ('complete', 'failed'):
                break
        self.assertEqual(job.status, 'complete', job.error)

    def test_dates_sorted_cumulative_shortage_and_no_stock_mutation(self):
        job = self.job('20260902,FG,80\n20260901,FG,80')
        self.finish(job)
        results = list(job.rows.values_list('payload', flat=True))
        self.assertEqual([r['need_date'] for r in results], ['2026-09-01', '2026-09-02'])
        self.assertEqual([r['items'][0]['stock_qty'] for r in results], [100, 20])
        self.assertEqual([r['items'][0]['shortage'] for r in results], [0, 60])
        self.stock.refresh_from_db()
        self.assertEqual(self.stock.quantity, 100)
        self.assertFalse(Demand.objects.exists())

    def test_snapshot_queries_constant_and_nested_bom_matches_original(self):
        semi = Product.objects.create(part_no='SEMI', part_name='Semi')
        BOMItem.objects.create(product=self.product, seq=2, child_part_no='SEMI', child_part_name='Semi', required_qty=2)
        BOMItem.objects.create(product=semi, child_part_no='RAW', child_part_name='Raw material', required_qty=3)
        plans = [{'part_no': 'FG', 'qty': 2, 'need_date': ''} for _ in range(90)]
        with CaptureQueriesContext(connection) as queries:
            templates, stocks, _ = build_snapshot(plans)
        self.assertEqual(len(queries), 4)
        _, _, old_items, old_structured = _calculate_bom_requirements('FG', 2)
        self.assertEqual(templates['FG']['items'][0]['required_qty'] * 2, old_items[0]['required_qty'])
        self.assertEqual([i['required_qty'] * 2 for i in templates['FG']['structured_items']], [i['required_qty'] for i in old_structured])
        job = self.job('20260901,FG,20')
        self.finish(job)
        result = job.rows.get().payload
        self.assertEqual(result['items'][0]['shortage'], 40)
        self.assertEqual(sum(i.get('shortage', 0) for i in result['structured_items']), 40)

    def test_ninety_days_resume_and_repeat_step_do_not_double_allocate(self):
        rows = '\n'.join(f'{date(2026, 1, 1) + timedelta(days=i)},FG,2' for i in range(90))
        job = self.job(rows)
        stale_copy = BOMCalculationJob.objects.get(pk=job.pk)
        advance_job(job)
        advance_job(stale_copy)
        job.refresh_from_db()
        self.assertEqual(job.completed, 25)
        self.assertEqual(job.remaining['RAW'], 50)
        self.finish(job)
        advance_job(job)
        self.assertEqual(job.rows.filter(ready=True).count(), 90)
        self.assertEqual(sum(r['items'][0]['shortage'] for r in job.rows.values_list('payload', flat=True)), 80)

    def test_failed_chunk_rolls_back_results_and_can_retry(self):
        job = self.job()
        with patch.object(BOMCalculationRow.objects, 'bulk_update', side_effect=RuntimeError('test failure')):
            with self.assertLogs('material.bom_jobs', level='ERROR'):
                advance_job(job)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertEqual(job.completed, 0)
        self.assertFalse(job.rows.filter(ready=True).exists())
        self.assertEqual(job.remaining['RAW'], 100)
        self.finish(job)

    def test_live_lease_waits_stale_lease_can_resume(self):
        job = self.job()
        BOMCalculationJob.objects.filter(pk=job.pk).update(status='running')
        job.refresh_from_db()
        advance_job(job)
        self.assertFalse(job.rows.filter(ready=True).exists())
        BOMCalculationJob.objects.filter(pk=job.pk).update(updated_at=timezone.now() - timedelta(minutes=3))
        job.refresh_from_db()
        self.finish(job)

    def test_cycle_rejected_without_partial_job(self):
        BOMItem.objects.create(product=self.product, seq=2, child_part_no='FG', child_part_name='Cycle', required_qty=1)
        with self.assertRaisesMessage(ValueError, '순환 참조'):
            self.job()
        self.assertFalse(BOMCalculationJob.objects.exists())

    def test_pivot_xlsx_and_quoted_csv(self):
        wb = openpyxl.Workbook()
        wb.active.append(['품번', '2026-09-02(수)', date(2026, 9, 1)])
        wb.active.append(['FG', 2, 3])
        stream = io.BytesIO()
        wb.save(stream)
        plans = parse_plans(SimpleUploadedFile('plan.xlsx', stream.getvalue()))
        self.assertEqual([(p['need_date'], p['qty']) for p in plans], [('2026-09-01', 3), ('2026-09-02', 2)])
        self.assertEqual(parse_plans(csv_file('20260901,FG,"1,200"'))[0]['qty'], 1200)
        with self.assertRaisesMessage(ValueError, '날짜'):
            parse_plans(csv_file('20260230,FG,2'))

    def test_permissions_ownership_csrf_and_pending_exports(self):
        job = self.job()
        key = f'job:{job.pk}'
        other = User.objects.create_user(username='other-admin', is_superuser=True)
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse('material:bom_calculate'), {'job': job.pk}).status_code, 404)
        self.assertEqual(self.client.post(reverse('material:bom_job_step', args=[job.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('material:bom_calc_batch_export'), {'session_key': key}).status_code, 404)
        self.assertEqual(self.client.post(reverse('material:bom_register_demand'), {'session_key': key}).status_code, 404)
        staff = User.objects.create_user(username='no-bom')
        self.client.force_login(staff)
        self.assertEqual(self.client.post(reverse('material:bom_job_start'), {'calc_file': csv_file('20260901,FG,2')}).status_code, 302)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('material:bom_calc_batch_export'), {'session_key': key}).status_code, 302)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.user)
        self.assertEqual(csrf_client.post(reverse('material:bom_job_step', args=[job.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse('material:bom_job_step', args=[job.pk])).status_code, 405)

    def test_pagination_all_metadata_exports_and_demand_use_entire_job(self):
        job = self.job('\n'.join(f'20260901,FG,1' for _ in range(65)) + '\n20260902,UNKNOWN,5')
        self.finish(job)
        url = reverse('material:bom_calculate')
        page = self.client.get(url, {'job': job.pk})
        self.assertEqual(page.status_code, 200)
        self.assertEqual(len(page.context['batch_results']), 50)
        self.assertEqual(page.context['batch_total'], 66)
        self.assertEqual(page.context['batch_metadata']['missing'], ['UNKNOWN'])
        self.assertEqual(page.context['batch_metadata']['vendors']['BOM vendor'], 65)
        self.assertEqual(page.content.count(b'<tbody class="batch-group"'), 50)
        missing = self.client.get(url, {'job': job.pk, 'filter': 'missing'})
        self.assertEqual(len(missing.context['batch_results']), 1)
        key = f'job:{job.pk}'
        for endpoint in ('bom_calc_batch_export', 'bom_calc_demand_export'):
            response = self.client.get(reverse('material:' + endpoint), {'session_key': key})
            self.assertEqual(response.status_code, 200)
            workbook = openpyxl.load_workbook(io.BytesIO(response.content))
            self.assertEqual(workbook.active.max_row, 67 if endpoint == 'bom_calc_batch_export' else 2)
            if endpoint == 'bom_calc_batch_export':
                self.assertEqual(workbook.active.cell(1, 8).value, '배분 전 가용재고')
        response = self.client.get(reverse('material:bom_calc_batch_export'), {'session_key': key, 'mode': 'structured'})
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.active.max_row, 132)
        response = self.client.post(reverse('material:bom_register_demand'), {'session_key': key})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Demand.objects.get(part=self.part).quantity, 65)

    def test_upload_replaces_large_legacy_session_and_progress_endpoint(self):
        session = self.client.session
        session['batch_calc_old'] = [{'part_no': 'FG'}]
        session.save()
        response = self.client.post(reverse('material:bom_job_start'), {'calc_file': csv_file('20260901,FG,5')})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('batch_calc_old', self.client.session)
        job = BOMCalculationJob.objects.get()
        page = self.client.get(response.json()['url'])
        self.assertContains(page, 'bomJobProgress')
        progress = self.client.post(reverse('material:bom_job_step', args=[job.pk])).json()
        self.assertEqual(progress['completed'], progress['total'])
        self.assertEqual(progress['status'], 'complete')

    def test_missing_date_allocated_last_and_shared_material_between_products(self):
        other = Product.objects.create(part_no='FG2', part_name='Second')
        BOMItem.objects.create(product=other, child_part_no='RAW', child_part_name='Raw', required_qty=2)
        job = self.job(',FG,50\n20260901,FG2,40')
        self.finish(job)
        self.assertEqual([p['items'][0]['shortage'] for p in job.rows.values_list('payload', flat=True)], [0, 30])

    def test_snapshot_is_stable_after_stock_and_bom_change(self):
        job = self.job()
        MaterialStock.objects.filter(pk=self.stock.pk).update(quantity=1000)
        BOMItem.objects.filter(product=self.product).update(required_qty=10)
        self.finish(job)
        self.assertEqual([p['items'][0]['shortage'] for p in job.rows.values_list('payload', flat=True)], [0, 60])

    def test_legacy_session_export_and_normal_form_submission(self):
        session = self.client.session
        _, name, items, structured = _calculate_bom_requirements('FG', 5)
        session['batch_calc_legacy'] = [{'part_no': 'FG', 'part_name': name, 'qty': 5,
                                       'need_date': '2026-09-01', 'items': items, 'structured_items': structured}]
        session.save()
        response = self.client.get(reverse('material:bom_calc_batch_export'), {'session_key': 'legacy'})
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.active.cell(1, 8).value, '현재고')
        response = self.client.post(reverse('material:bom_calculate'), {'calc_type': 'batch', 'calc_file': csv_file('20260901,FG,5')})
        self.assertEqual(response.status_code, 302)
        self.assertIn('?job=', response.url)
