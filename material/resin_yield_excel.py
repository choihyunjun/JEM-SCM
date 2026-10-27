"""원단위 중량 관리 엑셀: 성형팀이 쓰던 '생산품목별 원단위 중량 관리기준 자료' 양식.

시트 1 '원단위 중량 관리기준': 품번별 표 + [수지별 소계] + 합계
시트 2 '원재료별 사용량', 시트 3 '확인 필요'
"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
RED = 'C00000'
FILL_HEAD = PatternFill('solid', start_color=RED, end_color=RED)
FILL_UNIT = PatternFill('solid', start_color='DDEBF7', end_color='DDEBF7')    # 원단위
FILL_LOSS = PatternFill('solid', start_color='FFF2CC', end_color='FFF2CC')    # 스크랩 발생량
FILL_RATE = PatternFill('solid', start_color='FCE4E4', end_color='FCE4E4')    # 스크랩률
FILL_STRIPE = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
FILL_TOTAL = PatternFill('solid', start_color='D6DCE4', end_color='D6DCE4')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(left=THIN, right=THIN, top=Side(style='double', color='404040'), bottom=Side(style='thin', color='404040'))

EA = '#,##0;-#,##0;"-"'
KG = '#,##0.00;-#,##0.00;"-"'
G = '#,##0.00;-#,##0.00;"-"'
PCT = '0.0%'


def _title(ws, last_col, title, note):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A2'] = note
    ws['A2'].font = Font(name=FONT, size=9, color='808080')
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 26


def _header(ws, row, headers):
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row, col, text)
        cell.fill = FILL_HEAD
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def _rate_font(rate, bold=False):
    if rate >= 0.5:
        return Font(name=FONT, size=10, bold=True, color=RED)
    if rate >= 0.3:
        return Font(name=FONT, size=10, bold=bold, color='ED7D31')
    return Font(name=FONT, size=10, bold=bold)


def _write_row(ws, row, values, formats, fills=None, stripe=False, bold=False, border=BORDER):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(name=FONT, size=10, bold=bold)
        cell.border = border
        if col in formats:
            cell.number_format = formats[col]
        if isinstance(value, (int, float)) or value in ('-', '포함'):
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(vertical='center')
        if fills and col in fills:
            cell.fill = fills[col]
        elif stripe:
            cell.fill = FILL_STRIPE


def _widths(ws, widths):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _page(ws):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '3:3'


def _avg(values):
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else None


def _resin_subtotals(lines):
    groups = {}
    for r in lines:
        g = groups.setdefault(r['resin'], {'lines': [], 'parts': set(), 'produced': 0.0})
        g['lines'].append(r)
        if r['part_no'] not in g['parts']:
            g['parts'].add(r['part_no'])
            g['produced'] += r['produced']
    return groups


def _subtotal_values(lines, part_count, produced):
    total = sum(r['total_kg'] for r in lines)
    loss = sum(r['loss_kg'] for r in lines)
    return [
        part_count, None, None, None,
        _avg([r['product_weight'] for r in lines]), _avg([r['runner_weight'] for r in lines]),
        _avg([r['unit_standard'] for r in lines]), produced,
        sum(r['good_kg'] for r in lines), sum(r['bad_kg'] + r['inspect_kg'] for r in lines),
        sum(r['runner_kg'] for r in lines), sum(r['scrap_kg'] for r in lines), total, loss,
        loss / total if total else 0,
    ]


def build_workbook(report, year, month):
    wb = openpyxl.Workbook()
    lines = report['lines']
    source = (f"※ 신재투입량, 스크랩 발생량 기준 | 원본: ERP 생산실적({report['fetched_at']} 조회) · 성형 마스터 · "
              f"3공장 창고→불량창고 이동(검사불량, 양품에서 차감)")

    # ── 시트 1: 원단위 중량 관리기준 ──
    ws = wb.active
    ws.title = '원단위 중량 관리기준'
    headers = ['NO', '품번', '품명', '수지명\n(재질)', '재료명', 'C/V', '제품중량\n(g/EA)', '런너중량\n(g/SHOT)',
               '원단위\n(g/EA)', '생산수\n(EA)', '양품량\n(kg)', '불량량\n(kg)', '런너량\n(kg)', '파지량\n(kg)',
               '총 수지사용량\n(kg)', '스크랩 발생량\n(kg)', '스크랩률\n(%)', '비고']
    _title(ws, len(headers), f'생산품목별 원단위 중량 관리기준 자료 ({year}년 {month}월)', source)
    _header(ws, 3, headers)
    formats = {6: '0', 7: G, 8: G, 9: G, 10: EA, 11: KG, 12: KG, 13: KG, 14: KG, 15: KG, 16: KG, 17: PCT}
    fills = {9: FILL_UNIT, 16: FILL_LOSS, 17: FILL_RATE}
    row = 4
    for i, r in enumerate(lines, 1):
        bom = r['source'] == 'BOM'
        _write_row(ws, row, [
            i, r['part_no'], r['part_name'], r['resin'], r['material_name'] or r['material_part_no'],
            '-' if bom else r['cavity'], '-' if bom else r['product_weight'], '-' if bom else r['runner_weight'],
            r['unit_standard'], r['produced'], r['good_kg'], r['bad_kg'] + r['inspect_kg'],
            '포함' if bom else r['runner_kg'], r['scrap_kg'], r['total_kg'], r['loss_kg'], r['loss_rate'],
            'BOM 기준 (런너 포함)' if bom else ('C/V 미입력' if r['no_cavity'] else ''),
        ], formats, fills)
        ws.cell(row, 17).font = _rate_font(r['loss_rate'])
        row += 1
    last_data_row = row - 1
    ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(last_data_row, 3)}'
    ws.freeze_panes = 'C4'

    # [수지별 소계]
    row += 1
    ws.cell(row, 1, '[ 수지별 소계 ]').font = Font(name=FONT, size=11, bold=True)
    row += 1
    _header(ws, row, ['NO', '수지명(재질)', '품목 수', '', '', '', '평균 제품중량\n(g/EA)', '평균 런너중량\n(g/SHOT)',
                      '평균 원단위\n(g/EA)', '생산수 합계\n(EA)', '양품량 합계\n(kg)', '불량량 합계\n(kg)',
                      '런너량 합계\n(kg)', '파지량 합계\n(kg)', '총 수지사용량\n(kg)', '스크랩 발생량\n(kg)',
                      '스크랩률\n(%)', ''])
    ws.row_dimensions[row].height = 46
    row += 1
    sub_formats = {3: EA, 7: G, 8: G, 9: G, 10: EA, 11: KG, 12: KG, 13: KG, 14: KG, 15: KG, 16: KG, 17: PCT}
    groups = _resin_subtotals(lines)
    for i, (resin, g) in enumerate(sorted(groups.items(), key=lambda kv: kv[0]), 1):
        values = _subtotal_values(g['lines'], len(g['parts']), g['produced'])
        _write_row(ws, row, [i, resin] + values + [''], sub_formats, stripe=i % 2 == 0)
        ws.cell(row, 2).alignment = Alignment(horizontal='center')
        ws.cell(row, 17).font = _rate_font(values[-1])
        row += 1
    if lines:
        produced = sum(g['produced'] for g in groups.values())
        parts = len({r['part_no'] for r in lines})
        values = _subtotal_values(lines, parts, produced)
        _write_row(ws, row, ['합  계', ''] + values + [''], sub_formats, bold=True, border=TOTAL_BORDER)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1).alignment = Alignment(horizontal='center')
        for col in range(1, len(headers) + 1):
            ws.cell(row, col).fill = FILL_TOTAL
        ws.cell(row, 17).font = _rate_font(values[-1], bold=True)
    _widths(ws, [6, 16, 30, 10, 30, 6, 10, 10, 10, 12, 12, 11, 11, 11, 13, 13, 10, 20])
    _page(ws)

    # ── 시트 2: 원재료별 사용량 ──
    ws = wb.create_sheet('원재료별 사용량')
    headers = ['NO', '재료품번', '재료명', '수지명\n(재질)', '품목 수', '생산수\n(EA)', '양품량\n(kg)', '불량량\n(kg)',
               '런너량\n(kg)', '파지량\n(kg)', '총 수지사용량\n(kg)', '스크랩 발생량\n(kg)', '스크랩률\n(%)']
    _title(ws, len(headers), f'원재료별 수지 사용량 ({year}년 {month}월)', source)
    _header(ws, 3, headers)
    formats = {5: EA, 6: EA, 7: KG, 8: KG, 9: KG, 10: KG, 11: KG, 12: KG, 13: PCT}
    fills = {12: FILL_LOSS, 13: FILL_RATE}
    row = 4
    for i, s in enumerate(report['by_material'], 1):
        _write_row(ws, row, [i, s['key'], s['material_name'], s['resin'], s['part_count'], s['produced'],
                             s['good_kg'], s['bad_kg'] + s['inspect_kg'], s['runner_kg'], s['scrap_kg'],
                             s['total_kg'], s['loss_kg'], s['loss_rate']], formats, fills)
        ws.cell(row, 13).font = _rate_font(s['loss_rate'])
        row += 1
    ws.auto_filter.ref = f'A3:M{max(row - 1, 3)}'
    t = report['totals']
    if t:
        _write_row(ws, row, ['합  계', '', '', '', t['part_count'], t['produced'], t['good_kg'],
                             t['bad_kg'] + t['inspect_kg'], t['runner_kg'], t['scrap_kg'], t['total_kg'],
                             t['loss_kg'], t['loss_rate']], formats, bold=True, border=TOTAL_BORDER)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row, 1).alignment = Alignment(horizontal='center')
        for col in range(1, len(headers) + 1):
            ws.cell(row, col).fill = FILL_TOTAL
        ws.cell(row, 13).font = _rate_font(t['loss_rate'], bold=True)
    ws.freeze_panes = 'C4'
    _widths(ws, [6, 12, 34, 10, 8, 12, 12, 11, 11, 11, 13, 13, 10])
    _page(ws)

    # ── 시트 3: 확인 필요 ──
    ws = wb.create_sheet('확인 필요')
    headers = ['구분', '품번', '품명', '설비', '실적일', '실적번호', '내용', '처리']
    _title(ws, len(headers), f'원단위 중량 관리 확인 필요 ({year}년 {month}월)',
           '성형 마스터 미등록·재질 미분류·파지 비고 형식 등 집계 정확도에 영향을 주는 항목')
    _header(ws, 3, headers)
    rows = (
        [['파지 비고 형식', i['part_no'], i['part_name'], i['equip'], i['date'], i['wr_no'],
          f"비고 '{i['remark']}': {i['issue']}", f"{i['grams']:,}g 반영"] for i in report['issues']]
        + [['성형 마스터 미등록', u['part_no'], u['part_name'], '', '', '',
            f"생산 {u['produced']:,.0f}EA / 파지 {u['scrap_kg']:,.2f}kg (BOM도 없음)", '성형 마스터 등록 필요 (집계 제외)']
           for u in report['unregistered']]
        + [['BOM 기준 계산', p, '', '', '', '', '제품·런너 중량 구분 없음 (런너가 양품·불량에 포함)',
            '성형 마스터 등록 권장'] for p in report.get('bom_parts', [])]
        + [['재질 미분류', g['key'], g['material_name'], '', '', '', f"총 사용량 {g['total_kg']:,.1f}kg",
            '관리자 > 수지 재질에 등록'] for g in report.get('unclassified', [])]
        + [['이전 생산분 검사불량', d['part_no'], d['part_name'], '', '', '',
            f"{d['material_part_no']} 검사불량 {d['defect_qty']:,.0f}EA = {d['kg']:,.3f}kg (이 달 생산 없음)",
            '품번별 표·합계에서 제외'] for d in report.get('prior_defects', [])]
        + [['C/V 미입력', r['part_no'], r['part_name'], '', '', '', r['material_part_no'], '런너량 0으로 계산']
           for r in report['no_cavity']]
    )
    for i, values in enumerate(rows, 4):
        _write_row(ws, i, values, {}, stripe=i % 2 == 1)
    ws.auto_filter.ref = f'A3:H{max(len(rows) + 3, 3)}'
    ws.freeze_panes = 'A4'
    _widths(ws, [16, 16, 30, 8, 10, 15, 50, 30])
    _page(ws)
    return wb
