from datetime import date, datetime, timezone as datetime_timezone
from io import BytesIO
import re

from django.contrib.auth.models import User
from django.test import TestCase, RequestFactory
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Incoming, Part, UserProfile, Vendor
from .views import _incoming_history_queryset


class IncomingHistoryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create(username='history-admin', is_superuser=True)
        self.vendor = Vendor.objects.create(code='HIST-A', name='예시 협력사')
        self.other_vendor = Vendor.objects.create(code='HIST-B', name='다른 협력사')
        self.part = Part.objects.create(part_no='000123', part_name='입고 이력 예시', vendor=self.vendor)
        self.other_part = Part.objects.create(part_no='OTHER', part_name='다른 제품', vendor=self.other_vendor)
        self.client.force_login(self.user)

    def incoming(self, total, good, day=14, part=None):
        return Incoming.objects.create(part=part or self.part, in_date=date(2026, 10, day),
                                       quantity=total, confirmed_qty=good)

    def exported(self, params=None):
        response = self.client.get(reverse('incoming_export'), params or {})
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content)).active

    def test_good_bad_and_unconfirmed_quantities_match_screen_and_export(self):
        rows = [(100, 100, 0), (100, 80, 20), (1, 0, 1), (100, None, None)]
        for day, (total, good, bad) in enumerate(rows, 1):
            incoming = self.incoming(total, good, day=day)
            self.assertEqual(incoming.defective_qty, bad)
        response = self.client.get(reverse('incoming_list'))
        # Column semantics must remain intact when presentation classes change.
        headers = re.findall(r'<th\b[^>]*>\s*([^<]+?)\s*</th>', response.content.decode())
        self.assertIn('입고수량', headers)
        self.assertIn('불량수량', headers)
        actual = [(i.quantity, i.confirmed_qty, i.defective_qty) for i in response.context['incomings']]
        ws = self.exported()
        exported = [tuple(row) for row in ws.iter_rows(min_row=6, min_col=5, max_col=7, values_only=True)]
        self.assertEqual(actual, list(reversed(rows)))
        self.assertEqual(exported, actual)

    def test_export_preserves_numeric_dates_local_time_and_text_identifiers(self):
        row = self.incoming(12000, 11000)
        stamp = datetime(2026, 10, 14, 1, 23, tzinfo=datetime_timezone.utc)
        Incoming.objects.filter(pk=row.pk).update(created_at=stamp)
        ws = self.exported()
        self.assertEqual(ws['C6'].value, '000123')
        self.assertEqual(ws['C6'].data_type, 's')
        self.assertEqual(ws['A6'].value.date(), row.in_date)
        self.assertEqual(ws['H6'].value, timezone.localtime(stamp).replace(tzinfo=None))
        self.assertEqual(ws['E6'].value, 12000)
        self.assertEqual(ws['G6'].value, 1000)
        self.assertEqual(ws['G6'].number_format, '#,##0')

    def test_vendor_period_and_search_filters_match_between_list_and_export(self):
        wanted = self.incoming(100, 80, day=10)
        self.incoming(500, 500, day=11, part=self.other_part)
        self.incoming(200, 190, day=12)
        params = {'vendor_id': self.vendor.pk, 'start_date': '2026-10-10',
                  'end_date': '2026-10-11', 'q': '000123'}
        response = self.client.get(reverse('incoming_list'), params)
        self.assertEqual(list(response.context['incomings'].values_list('pk', flat=True)), [wanted.pk])
        ws = self.exported(params)
        self.assertEqual(ws.max_row, 6)
        self.assertEqual(ws['E6'].value, 100)
        # 시작일 또는 종료일만 입력해도 두 경로에 동일하게 반영한다.
        for params in ({'start_date': '2026-10-12'}, {'end_date': '2026-10-10'}):
            response = self.client.get(reverse('incoming_list'), params)
            ws = self.exported(params)
            self.assertEqual(response.context['incomings'].count(), ws.max_row - 5)

    def test_supplier_scope_cannot_be_overridden_by_vendor_filter(self):
        self.incoming(100, 80)
        self.incoming(100, 100, part=self.other_part)
        user = User.objects.create(username='history-supplier')
        UserProfile.objects.filter(user=user).update(role='VENDOR', org=self.vendor.organization)
        request = RequestFactory().get('/', {'vendor_id': self.other_vendor.pk})
        request.user = User.objects.get(pk=user.pk)
        self.assertEqual(list(_incoming_history_queryset(request).values_list('part_id', flat=True)), [self.part.pk])

    def test_download_has_header_stripes_filter_panes_and_defect_highlighting(self):
        self.incoming(100, 80, day=10)
        self.incoming(200, 200, day=11)
        ws = self.exported()
        self.assertEqual(ws.title, '입고 이력')
        self.assertEqual(ws.freeze_panes, 'D6')
        self.assertEqual(ws.auto_filter.ref, 'A5:H7')
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws['A5'].fill.fgColor.rgb, '00243B53')
        self.assertNotEqual(ws['A6'].fill.fgColor.rgb, ws['A7'].fill.fgColor.rgb)
        self.assertEqual(len(ws.conditional_formatting), 1)
        self.assertEqual(ws.page_setup.orientation, 'landscape')

    def test_empty_download_and_missing_vendor_still_work(self):
        ws = self.exported({'q': 'NOT-FOUND'})
        self.assertEqual(ws.auto_filter.ref, 'A5:H5')
        self.assertEqual(ws['H2'].value, '0건')
        part = Part.objects.create(part_no='NO-VENDOR', part_name='업체 미지정')
        self.incoming(10, 10, part=part)
        ws = self.exported()
        self.assertIsNone(ws['B6'].value)

    def test_formula_like_names_are_exported_as_literal_text(self):
        self.part.part_name = '=1+1'
        self.part.save(update_fields=['part_name'])
        self.incoming(1, 0)
        ws = self.exported()
        self.assertEqual(ws['D6'].value, '=1+1')
        self.assertEqual(ws['D6'].data_type, 's')
