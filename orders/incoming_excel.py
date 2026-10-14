"""입고 이력 다운로드. 기존 openpyxl 기반 내보내기에서 사용하는 서식."""
import math
import unicodedata
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from django.utils import timezone


def _wrapped_lines(value, width):
    lines = str(value or '').split('\n')
    return sum(max(1, math.ceil(sum(
        2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1 for char in line
    ) / width)) for line in lines)


def incoming_workbook(incomings):
    wb = Workbook()
    ws = wb.active
    ws.title = '입고 이력'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '243B53'
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 8
    ws['A2'] = '입고 처리 이력'
    ws['A2'].font = Font(name='맑은 고딕', size=16, bold=True, color='243B53')
    ws['A2'].alignment = Alignment(vertical='center')
    ws['A3'] = '입고수량 = 확정 양품 / 불량수량 = 납품수량 − 입고수량 / 미확정 수량은 공란'
    ws['A3'].font = Font(name='맑은 고딕', size=10, color='62748A')
    ws['A3'].alignment = Alignment(vertical='center')

    headers = ['입고일자', '협력사', '품번', '품명', '납품수량', '입고수량', '불량수량', '처리일시']
    widths = [15, 24, 25, 46, 15, 15, 15, 23]
    for col, (label, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(5, col, label)
        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(right=Side(style='thin', color='FFFFFF'))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[5].height = 30

    text_font = Font(name='맑은 고딕', size=11, color='243B53')
    row_line = Border(bottom=Side(style='hair', color='DDE5ED'))
    count = 0
    for row, item in enumerate(incomings, 6):
        count += 1
        processed_at = timezone.localtime(item.created_at).replace(tzinfo=None)
        values = [item.in_date, item.part.vendor.name if item.part.vendor else '',
                  item.part.part_no, item.part.part_name, item.quantity,
                  item.confirmed_qty, item.defective_qty, processed_at]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.font = text_font
            cell.border = row_line
            cell.fill = PatternFill('solid', fgColor='F3F6FA' if row % 2 else 'FFFFFF')
            cell.alignment = Alignment(horizontal='right' if col in (1, 5, 6, 7, 8) else 'left',
                                       vertical='center', indent=1, wrap_text=col in (2, 3, 4))
            if col in (2, 3, 4):
                # 품번/업체명은 '='로 시작해도 Excel 수식이 아닌 원문 문자열이다.
                cell.data_type = 's'
            elif col in (5, 6, 7):
                cell.number_format = '#,##0'
            elif col == 1:
                cell.number_format = 'yyyy-mm-dd'
            elif col == 8:
                cell.number_format = 'yyyy-mm-dd hh:mm'
        # 긴 품명도 인쇄할 때 가려지지 않도록 내용 길이에 맞춰 행 높이를 잡는다.
        lines = max(_wrapped_lines(values[col - 1], widths[col - 1] - 2) for col in (2, 3, 4))
        ws.row_dimensions[row].height = max(27, lines * 17 + 6)

    last_row = max(5, count + 5)
    ws['H2'] = f'{count:,}건'
    ws['H2'].font = Font(name='맑은 고딕', size=11, color='62748A')
    ws['H2'].alignment = Alignment(horizontal='right', vertical='center')
    ws.auto_filter.ref = f'A5:H{last_row}'
    ws.freeze_panes = 'D6'
    if count:
        ws.conditional_formatting.add(f'G6:G{last_row}', CellIsRule(
            operator='greaterThan', formula=['0'],
            fill=PatternFill('solid', fgColor='FDEBEC'),
            font=Font(name='맑은 고딕', size=11, bold=True, color='B42332'),
        ))
    ws.print_title_rows = '1:5'
    ws.print_area = f'A1:H{last_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    ws.oddFooter.center.text = '&P / &N'
    return wb
